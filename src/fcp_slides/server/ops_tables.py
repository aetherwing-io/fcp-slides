"""Table operation handlers — add, set, style, row, header, merge, remove."""

from __future__ import annotations

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

from fcp_core import OpResult, ParsedOp

from fcp_slides.lib.colors import parse_color, to_rgb
from fcp_slides.lib.units import parse_length
from fcp_slides.model.refs import ShapeRef
from fcp_slides.server.resolvers import (
    SlidesOpContext,
    extract_position,
    require_active_slide,
    resolve_shape_on_slide,
)


_DEFAULT_TABLE_LEFT = Inches(1)
_DEFAULT_TABLE_TOP = Inches(2)
_DEFAULT_TABLE_WIDTH = Inches(8)
_DEFAULT_TABLE_HEIGHT = Inches(4)


def op_table(op: ParsedOp, ctx: SlidesOpContext) -> OpResult:
    """Dispatch table sub-commands."""
    if not op.positionals:
        return OpResult(success=False, message="Usage: table add|set|style|row|header|merge|remove ...")

    action = op.positionals[0].lower()
    rest = op.positionals[1:]

    dispatch = {
        "add": _table_add,
        "set": _table_set,
        "style": _table_style,
        "row": _table_row,
        "header": _table_header,
        "merge": _table_merge,
        "remove": _table_remove,
    }

    handler = dispatch.get(action)
    if handler is None:
        return OpResult(
            success=False,
            message=f"Unknown table action: {action!r}. Use: add, set, style, row, header, merge, remove",
        )

    return handler(rest, op.params, ctx)


def _table_add(
    args: list[str], params: dict[str, str], ctx: SlidesOpContext
) -> OpResult:
    """Add a table to the active slide.

    Usage: table add ROWS COLS [label:NAME] [x:POS] [y:POS] [w:SIZE] [h:SIZE]
    """
    result = require_active_slide(ctx)
    if isinstance(result, str):
        return OpResult(success=False, message=result)
    slide, slide_idx = result

    if len(args) < 2:
        return OpResult(success=False, message="Usage: table add ROWS COLS [label:NAME]")

    try:
        rows = int(args[0])
        cols = int(args[1])
    except ValueError:
        return OpResult(success=False, message="ROWS and COLS must be integers")

    if rows < 1 or cols < 1:
        return OpResult(success=False, message="ROWS and COLS must be >= 1")

    pos = extract_position(params)
    left = pos.get("left", _DEFAULT_TABLE_LEFT)
    top = pos.get("top", _DEFAULT_TABLE_TOP)
    width = pos.get("width", _DEFAULT_TABLE_WIDTH)
    height = pos.get("height", _DEFAULT_TABLE_HEIGHT)

    shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = shape.table

    label = params.get("label", "")
    if label:
        ref = ShapeRef(
            label=label,
            slide_idx=slide_idx,
            shape_id=shape.shape_id,
            shape_type="table",
        )
        ctx.index.add_shape_label(label, ref)
    else:
        ctx.index.rebuild(ctx.model)

    return OpResult(
        success=True,
        message=f"Table {rows}x{cols} added" + (f" label:{label}" if label else ""),
        prefix="+",
    )


def _resolve_table_shape(ref: str, ctx: SlidesOpContext):
    """Find a table shape by label on the active slide."""
    active = require_active_slide(ctx)
    if isinstance(active, str):
        return None, None, active
    slide, slide_idx = active

    shape = resolve_shape_on_slide(ref, slide, slide_idx, ctx)
    if shape is None:
        return None, None, f"Shape not found: {ref!r}"

    if not shape.has_table:
        return None, None, f"Shape '{ref}' is not a table"

    return shape, shape.table, None


def _table_set(
    args: list[str], params: dict[str, str], ctx: SlidesOpContext
) -> OpResult:
    """Set a cell value in a table.

    Usage: table set TABLE_REF ROW COL VALUE
    ROW, COL are 0-based.
    """
    if len(args) < 4:
        return OpResult(success=False, message='Usage: table set TABLE_REF ROW COL "VALUE"')

    table_ref = args[0]
    shape, table, err = _resolve_table_shape(table_ref, ctx)
    if err:
        return OpResult(success=False, message=err)

    try:
        row = int(args[1])
        col = int(args[2])
    except ValueError:
        return OpResult(success=False, message="ROW and COL must be integers")

    value = args[3]

    if row < 0 or row >= len(table.rows):
        return OpResult(success=False, message=f"Row {row} out of range (0-{len(table.rows) - 1})")
    if col < 0 or col >= len(table.columns):
        return OpResult(success=False, message=f"Col {col} out of range (0-{len(table.columns) - 1})")

    table.cell(row, col).text = value

    return OpResult(success=True, message=f"Table cell ({row},{col}) = \"{value[:30]}\"", prefix="*")


def _table_style(
    args: list[str], params: dict[str, str], ctx: SlidesOpContext
) -> OpResult:
    """Style table cells.

    Usage: table style TABLE_REF ROW_RANGE,COL_RANGE [bold] [fill:#HEX] [color:#HEX] [size:N]
    Example: table style metrics 0,0:0,3 bold fill:#4472C4 color:#FFFFFF
    """
    if len(args) < 2:
        return OpResult(success=False, message="Usage: table style TABLE_REF CELL_RANGE [bold] [fill:#HEX] [color:#HEX]")

    table_ref = args[0]
    shape, table, err = _resolve_table_shape(table_ref, ctx)
    if err:
        return OpResult(success=False, message=err)

    cell_range = args[1]
    cells = _parse_cell_range(cell_range, table)
    if isinstance(cells, str):
        return OpResult(success=False, message=cells)

    # Collect flags from remaining args
    flags = set()
    for a in args[2:]:
        if a.lower() in ("bold", "italic", "underline"):
            flags.add(a.lower())

    count = 0
    for row, col in cells:
        cell = table.cell(row, col)

        # Apply fill
        if "fill" in params:
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = to_rgb(params["fill"])

        # Apply text formatting
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                if "bold" in flags:
                    run.font.bold = True
                if "italic" in flags:
                    run.font.italic = True
                if "underline" in flags:
                    run.font.underline = True
                if "color" in params:
                    run.font.color.rgb = to_rgb(params["color"])
                if "size" in params:
                    run.font.size = Pt(float(params["size"]))
                if "font" in params:
                    run.font.name = params["font"]

        count += 1

    return OpResult(success=True, message=f"Styled {count} table cells", prefix="*")


def _parse_cell_range(spec: str, table) -> list[tuple[int, int]] | str:
    """Parse a cell range spec like '0,0:0,3' or '0' (whole row).

    Returns list of (row, col) tuples or error string.
    """
    num_rows = len(table.rows)
    num_cols = len(table.columns)

    if ":" in spec:
        # Range: start_row,start_col:end_row,end_col
        parts = spec.split(":")
        if len(parts) != 2:
            return f"Invalid range: {spec!r}"
        try:
            sr, sc = (int(x) for x in parts[0].split(","))
            er, ec = (int(x) for x in parts[1].split(","))
        except ValueError:
            return f"Invalid range: {spec!r}"

        cells = []
        for r in range(sr, er + 1):
            for c in range(sc, ec + 1):
                if 0 <= r < num_rows and 0 <= c < num_cols:
                    cells.append((r, c))
        return cells

    if "," in spec:
        # Single cell: row,col
        try:
            r, c = (int(x) for x in spec.split(","))
        except ValueError:
            return f"Invalid cell: {spec!r}"
        if 0 <= r < num_rows and 0 <= c < num_cols:
            return [(r, c)]
        return f"Cell ({r},{c}) out of range"

    # Whole row
    try:
        r = int(spec)
    except ValueError:
        return f"Invalid range: {spec!r}"
    if 0 <= r < num_rows:
        return [(r, c) for c in range(num_cols)]
    return f"Row {r} out of range"


def _table_row(
    args: list[str], params: dict[str, str], ctx: SlidesOpContext
) -> OpResult:
    """Set an entire row of values.

    Usage: table row TABLE_REF ROW_IDX VALUE1 VALUE2 ...
    """
    if len(args) < 3:
        return OpResult(success=False, message='Usage: table row TABLE_REF ROW_IDX "V1" "V2" ...')

    table_ref = args[0]
    shape, table, err = _resolve_table_shape(table_ref, ctx)
    if err:
        return OpResult(success=False, message=err)

    try:
        row_idx = int(args[1])
    except ValueError:
        return OpResult(success=False, message="ROW_IDX must be an integer")

    if row_idx < 0 or row_idx >= len(table.rows):
        return OpResult(success=False, message=f"Row {row_idx} out of range (0-{len(table.rows) - 1})")

    values = args[2:]
    num_cols = len(table.columns)
    for col, val in enumerate(values[:num_cols]):
        table.cell(row_idx, col).text = val

    return OpResult(
        success=True,
        message=f"Row {row_idx} set ({len(values[:num_cols])} values)",
        prefix="*",
    )


def _table_header(
    args: list[str], params: dict[str, str], ctx: SlidesOpContext
) -> OpResult:
    """Set header row values (row 0).

    Usage: table header TABLE_REF VALUE1 VALUE2 ...
    """
    if len(args) < 2:
        return OpResult(success=False, message='Usage: table header TABLE_REF "H1" "H2" ...')

    table_ref = args[0]
    shape, table, err = _resolve_table_shape(table_ref, ctx)
    if err:
        return OpResult(success=False, message=err)

    values = args[1:]
    num_cols = len(table.columns)
    for col, val in enumerate(values[:num_cols]):
        table.cell(0, col).text = val

    return OpResult(
        success=True,
        message=f"Header set ({len(values[:num_cols])} columns)",
        prefix="*",
    )


def _table_merge(
    args: list[str], params: dict[str, str], ctx: SlidesOpContext
) -> OpResult:
    """Merge table cells.

    Usage: table merge TABLE_REF START_ROW,START_COL:END_ROW,END_COL
    """
    if len(args) < 2:
        return OpResult(success=False, message="Usage: table merge TABLE_REF ROW,COL:ROW,COL")

    table_ref = args[0]
    shape, table, err = _resolve_table_shape(table_ref, ctx)
    if err:
        return OpResult(success=False, message=err)

    range_spec = args[1]
    if ":" not in range_spec:
        return OpResult(success=False, message="Merge range must be START_ROW,START_COL:END_ROW,END_COL")

    parts = range_spec.split(":")
    try:
        sr, sc = (int(x) for x in parts[0].split(","))
        er, ec = (int(x) for x in parts[1].split(","))
    except ValueError:
        return OpResult(success=False, message=f"Invalid range: {range_spec!r}")

    start_cell = table.cell(sr, sc)
    end_cell = table.cell(er, ec)
    start_cell.merge(end_cell)

    return OpResult(success=True, message=f"Merged cells ({sr},{sc}):({er},{ec})", prefix="*")


def _table_remove(
    args: list[str], params: dict[str, str], ctx: SlidesOpContext
) -> OpResult:
    """Remove a table (shape) from the slide."""
    if not args:
        return OpResult(success=False, message="Usage: table remove TABLE_REF")

    table_ref = args[0]
    active = require_active_slide(ctx)
    if isinstance(active, str):
        return OpResult(success=False, message=active)
    slide, slide_idx = active

    shape = resolve_shape_on_slide(table_ref, slide, slide_idx, ctx)
    if shape is None:
        return OpResult(success=False, message=f"Shape not found: {table_ref!r}")

    sp = shape._element
    sp.getparent().remove(sp)
    ctx.index.rebuild(ctx.model)

    return OpResult(success=True, message=f"Table '{table_ref}' removed", prefix="-")


HANDLERS: dict[str, callable] = {
    "table": op_table,
}
